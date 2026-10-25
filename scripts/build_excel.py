#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
clientcard2excel — 展会客户 Excel 生成引擎
=========================================
将「客户线索 leads」+「联系方式 contacts」生成为标准化 Excel：

  · 客户明细 sheet：14 列（8 个基础字段 + 6 个联系方式列）
  · 统计汇总 sheet：意向分布 / 地区分布 / 重点跟进排序
  · 品牌紫表头、意向强度色块编码、冻结窗格、自动筛选

推荐用法（AI 驱动）：
  1. 编辑下方「DATA — 在此填充」区的 LEADS / CONTACTS
  2. 运行：python build_excel.py
  3. 生成的 Excel 落在当前工作目录：<TITLE>客户明细.xlsx

也可作为模块 import：
  from build_excel import generate_excel
  generate_excel(leads, contacts, output_path="out.xlsx", title="上海玩具展")

依赖：openpyxl  （pip install openpyxl）
"""

import os
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ===================== 基础配置 =====================
TITLE = "展会客户"            # 文件名前缀，例如「香港秋季电子展」「上海玩具展」
                           # 不同 TITLE → 不同文件名，各展会独立建表、互不覆盖
SHEET_DETAIL = "客户明细"
SHEET_SUMMARY = "统计汇总"

# ===================== 视觉规范 =====================
BRAND_PURPLE = "6B4E9E"     # 品牌紫（表头底色）
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
DETAIL_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="D9D2E9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 意向强度 → 背景色块（语义色）
INTENT_FILL = {
    "已下单":   "C8E6C9",   # 绿：已成交
    "已购样品": "FFE0B2",   # 橙：已付钱买样
    "强":       "F8BBD0",   # 粉红：强意向
    "中":       "FFF9C4",   # 浅黄：一般意向（低饱和）
    "低":       "ECEFF1",   # 灰：低意向
}
# 跟进优先级（数值越大越优先）
INTENT_PRIORITY = {"已下单": 5, "强": 4, "已购样品": 3, "中": 2, "低": 1, "待补充": 0}

# 14 列定义：(列名, 默认宽度)
COLUMNS = [
    ("序号", 8),
    ("日期", 12),
    ("国家/地区", 16),
    ("渠道/客户类型", 22),
    ("意向品类", 18),
    ("意向强度", 12),
    ("备注", 42),
    ("后续动作", 24),
    ("联系人/职务", 24),
    ("公司名称", 32),
    ("电话", 20),
    ("邮箱", 30),
    ("微信", 16),
    ("地址", 42),
]
LEAD_FIELDS = 8   # leads 元组前 8 项为基础字段
# contacts 字典：{ 序号(str): (联系人/职务, 公司, 电话, 邮箱, 微信, 地址) }

# ===================== 分类逻辑 =====================
def classify_intent(level):
    """把意向强度文本归一为色块 key。"""
    s = (level or "").strip()
    if "已下单" in s:
        return "已下单"
    if "样品" in s or "已购" in s:
        return "已购样品"
    if "强" in s:
        return "强"
    if "低" in s:
        return "低"
    return "中"

def region_of(s):
    """把国家/地区文本归类大区，用于统计。"""
    s = (s or "").strip()
    if "待确认" in s or not s:
        return "待确认"
    if "西班牙" in s and "语" in s:     # 西班牙语地区（待确认）特判
        return "待确认"
    asia = ["中国", "香港", "台湾", "日本", "韩国", "泰国", "印度", "新加坡", "马来",
            "印尼", "越南", "蒙古", "菲律宾", "巴基", "土耳其", "阿联酋", "迪拜",
            "以色列", "沙特", "伊朗", "中东"]
    europe = ["俄罗斯", "德国", "法国", "英国", "意大利", "西班牙", "葡萄牙", "波兰",
              "捷克", "克罗地亚", "挪威", "瑞典", "芬兰", "荷兰", "比利时", "瑞士",
              "奥地利", "希腊", "匈牙利", "保加利亚", "罗马尼亚", "乌克兰", "丹麦",
              "爱尔兰", "欧洲"]
    americas = ["美国", "加拿大", "墨西哥", "巴西", "阿根廷", "巴拉圭", "智利", "哥伦", "秘鲁", "美洲"]
    africa = ["南非", "埃及", "尼日", "肯尼", "摩洛", "突尼", "非洲"]
    oceania = ["澳大利亚", "新西兰", "澳洲", "大洋洲"]
    for kw in asia:
        if kw in s:
            return "亚洲/中东"
    for kw in europe:
        if kw in s:
            return "欧洲"
    for kw in americas:
        if kw in s:
            return "美洲"
    for kw in africa:
        if kw in s:
            return "非洲"
    for kw in oceania:
        if kw in s:
            return "大洋洲"
    return "其他"

# ===================== 生成逻辑 =====================
def _build_summary(ws, leads, contacts, title):
    widths = {"A": 22, "B": 14, "C": 34, "D": 18, "E": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    def title_cell(r, text):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name="微软雅黑", bold=True, size=12, color=BRAND_PURPLE)
        return c

    r = 1
    title_cell(r, f"{title} — 客户统计汇总"); r += 2

    c = ws.cell(row=r, column=1, value="线索总数")
    c.font = DETAIL_FONT
    ws.cell(row=r, column=2, value=len(leads)).font = DETAIL_FONT
    r += 2

    title_cell(r, "意向强度分布"); r += 1
    intent_counter = Counter(classify_intent(l[5]) for l in leads)
    for key in ["已下单", "强", "已购样品", "中", "低"]:
        if intent_counter.get(key):
            cell = ws.cell(row=r, column=1, value=key)
            cell.font = DETAIL_FONT
            cell.fill = PatternFill("solid", fgColor=INTENT_FILL[key])
            ws.cell(row=r, column=2, value=intent_counter[key]).font = DETAIL_FONT
            r += 1
    r += 1

    title_cell(r, "地区分布"); r += 1
    region_counter = Counter(region_of(l[2]) for l in leads)
    for key, cnt in region_counter.most_common():
        ws.cell(row=r, column=1, value=key).font = DETAIL_FONT
        ws.cell(row=r, column=2, value=cnt).font = DETAIL_FONT
        r += 1
    r += 1

    title_cell(r, "重点跟进排序（按意向优先级）"); r += 1
    hdr = ["序号", "国家/地区", "公司 / 联系人", "微信", "意向"]
    for ci, h in enumerate(hdr, start=1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = PatternFill("solid", fgColor=BRAND_PURPLE)
        c.alignment = Alignment(horizontal="center")
    r += 1
    ranked = sorted(leads, key=lambda l: INTENT_PRIORITY.get(classify_intent(l[5]), 0), reverse=True)
    for l in ranked:
        seq = str(l[0])
        contact = contacts.get(seq, ("", "", "", "", "", ""))
        comp = contact[1] or ""
        name = contact[0] or ""
        wx = contact[4] or ""
        vals = [seq, l[2], f"{comp} {name}".strip(), wx, l[5]]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = DETAIL_FONT
            cell.border = BORDER
        r += 1

def generate_excel(leads, contacts, output_path=None, title=None):
    """生成标准化客户 Excel。

    leads: list[tuple]，每项 8 字段：序号,日期,国家地区,渠道类型,意向品类,意向强度,备注,后续动作
    contacts: dict{序号: (联系人/职务,公司,电话,邮箱,微信,地址)}
    """
    title = title or TITLE
    if output_path is None:
        output_path = os.path.join(os.getcwd(), f"{title}客户明细.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_DETAIL

    # 表头
    for ci, (name, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=ci, value=name)
        c.fill = PatternFill("solid", fgColor=BRAND_PURPLE)
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width

    # 数据行
    for ri, lead in enumerate(leads, start=2):
        seq = str(lead[0])
        contact = contacts.get(seq, ("", "", "", "", "", ""))
        row = list(lead[:LEAD_FIELDS]) + list(contact[:6])
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = DETAIL_FONT
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = BORDER
        intent_key = classify_intent(row[5])
        if intent_key in INTENT_FILL:
            ic = ws.cell(row=ri, column=6)
            ic.fill = PatternFill("solid", fgColor=INTENT_FILL[intent_key])
            ic.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(leads) + 1}"

    ws2 = wb.create_sheet(SHEET_SUMMARY)
    _build_summary(ws2, leads, contacts, title)

    wb.save(output_path)
    return output_path


# ===================== 示例数据（替换为你提取的真实数据）=====================
if __name__ == "__main__":
    LEADS = [
        ("1", "2026-10-13", "中国香港", "贸易/采购", "AI毛绒玩具", "强（意向很强）",
         "现场拿了 2 只样品，老板亲自来", "确认订单与采购量"),
        ("2", "2026-10-13", "南非", "进口商", "待确认", "中",
         "有名片，做多国进口", "邮件跟进报价"),
    ]
    CONTACTS = {
        "1": ("张三（采购经理）", "示例贸易有限公司", "+852 1234 5678",
              "zhang@sample.com", "zhangsan_wx", "香港中环 Sample 大厦 12F"),
        "2": ("John Doe", "Sample Imports (Pty) Ltd", "+27 11 123 4567",
              "john@sample.co.za", "johndoe", "Johannesburg, South Africa"),
    }
    out = generate_excel(LEADS, CONTACTS)
    print("已生成 Excel:", out)
